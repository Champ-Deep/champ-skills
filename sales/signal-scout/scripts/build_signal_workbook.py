#!/usr/bin/env python3
"""Build the signal-scout prospecting workbook from findings.json + run_config.json.

Usage:
    python build_signal_workbook.py --data findings.json --config run_config.json --out list.xlsx

findings.json : array of company objects (see references/agent-prompt-template.md), each
                contact optionally carrying a "pitch" string added in Stage 3.
run_config.json : title, signal_criteria, compiled_date, pitch_context, brand_color,
                layout ("person"|"company"), market_sources, disclaimer.
"""
import argparse
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

GREY = "595959"
LINK_BLUE = "0563C1"
CAUTION_RED = "9C0006"
HEADER_DARK = "1F1F1F"
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def light_tint(hex_color: str) -> str:
    """A very light tint of the brand color for zebra striping."""
    r, g, b = (int(hex_color[i : i + 2], 16) for i in (0, 2, 4))
    mix = lambda c: int(c + (255 - c) * 0.92)
    return f"{mix(r):02X}{mix(g):02X}{mix(b):02X}"


def write_link(cell, url, label):
    if url and url.lower() not in ("", "not found", "n/a", "-", "—"):
        cell.value = label
        cell.hyperlink = url
        cell.font = Font(name="Arial", size=10, color=LINK_BLUE, underline="single")
    else:
        cell.value = "Not found"
        cell.font = Font(name="Arial", size=10, color=GREY, italic=True)


def base_cell(cell, value, *, wrap=True, center=False, caution=False):
    cell.value = value if value not in ("", None) else "—"
    cell.font = Font(
        name="Arial", size=10, color=CAUTION_RED if caution else "000000"
    )
    cell.alignment = Alignment(
        vertical="top", wrap_text=wrap, horizontal="center" if center else "left"
    )


PERSON_COLUMNS = [
    ("Rank", 6), ("Company", 26), ("Sector", 20), ("Signal Type", 16),
    ("Signal Specifics", 34), ("Signal Date", 12), ("Source", 10),
    ("Role", 12), ("Name & Title", 32), ("Person LinkedIn", 16),
    ("Person Findings", 44), ("Personalized Pitch", 60),
    ("Company LinkedIn", 16), ("Website", 22), ("HQ", 20),
    ("Public Email", 26), ("Public Phone", 20), ("Notes / Caveats", 44),
]


def build(data, cfg, out_path):
    brand = cfg.get("brand_color", "E8611A")
    stripe = light_tint(brand)
    layout = cfg.get("layout", "person")

    wb = Workbook()
    ws = wb.active
    ws.title = "Signal List"

    if layout == "company":
        roles = []
        for comp in data:
            for c in comp.get("contacts", []):
                if c.get("role") and c["role"] not in roles:
                    roles.append(c["role"])
        columns = [("Rank", 6), ("Company", 26), ("Sector", 20), ("Signal Type", 16),
                   ("Signal Specifics", 34), ("Signal Date", 12), ("Source", 10)]
        for r in roles:
            columns += [(r, 30), (f"{r} LinkedIn", 16), (f"{r} Pitch", 50)]
        columns += [("Company LinkedIn", 16), ("Website", 22), ("HQ", 20),
                    ("Public Email", 26), ("Public Phone", 20), ("Notes / Caveats", 44)]
    else:
        columns = PERSON_COLUMNS

    ncols = len(columns)
    last = get_column_letter(ncols)

    ws.merge_cells(f"A1:{last}1")
    t = ws["A1"]
    t.value = cfg.get("title", "Signal Scout — Prospect List")
    t.font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor=brand)
    t.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26

    ws.merge_cells(f"A2:{last}2")
    m = ws["A2"]
    m.value = (
        f"Signal criteria: {cfg.get('signal_criteria', '')} | Compiled: "
        f"{cfg.get('compiled_date', '')} | Pitch context: {cfg.get('pitch_context', '—')}. "
        f"{cfg.get('disclaimer', 'Public sources only — no guessed emails, phones, or LinkedIn profiles.')}"
    )
    m.font = Font(name="Arial", size=9, italic=True, color=GREY)
    m.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 30

    HDR = 3
    for i, (h, w) in enumerate(columns, 1):
        cell = ws.cell(row=HDR, column=i, value=h)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_DARK)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[HDR].height = 22

    data = sorted(data, key=lambda d: d.get("rank", 10**9))
    row = HDR + 1
    band = 0

    for ci, comp in enumerate(data, 1):
        sig = comp.get("signal", {}) or {}
        rank = comp.get("rank", ci)
        notes = comp.get("notes", "")
        caution = str(notes).upper().startswith("CAUTION")

        if layout == "company":
            rows_here = [row]
            vals = {
                1: rank, 2: comp.get("company"), 3: comp.get("sector"),
                4: sig.get("type"), 5: sig.get("specifics"), 6: sig.get("date"),
            }
            for col, v in vals.items():
                base_cell(ws.cell(row=row, column=col), v, center=col in (1, 6))
            write_link(ws.cell(row=row, column=7), sig.get("source_url"), "Source ↗")
            col = 8
            by_role = {c.get("role"): c for c in comp.get("contacts", [])}
            for r in roles:
                c = by_role.get(r, {})
                nm = f"{c.get('name', '—')}" + (f" — {c['title']}" if c.get("title") else "")
                base_cell(ws.cell(row=row, column=col), nm if c else "—")
                write_link(ws.cell(row=row, column=col + 1), c.get("linkedin"), "LinkedIn ↗")
                base_cell(ws.cell(row=row, column=col + 2), c.get("pitch", ""))
                col += 3
            write_link(ws.cell(row=row, column=col), comp.get("company_linkedin"), "Company ↗")
            write_link(ws.cell(row=row, column=col + 1), comp.get("website"),
                       (comp.get("website") or "").replace("https://", "").rstrip("/") or "Not found")
            base_cell(ws.cell(row=row, column=col + 2), comp.get("hq"))
            base_cell(ws.cell(row=row, column=col + 3), comp.get("public_email"))
            base_cell(ws.cell(row=row, column=col + 4), comp.get("public_phone"))
            base_cell(ws.cell(row=row, column=col + 5), notes, caution=caution)
            row += 1
        else:
            contacts = comp.get("contacts", []) or [{}]
            rows_here = list(range(row, row + len(contacts)))
            for c in contacts:
                base_cell(ws.cell(row=row, column=1), rank, center=True)
                base_cell(ws.cell(row=row, column=2), comp.get("company"))
                base_cell(ws.cell(row=row, column=3), comp.get("sector"))
                base_cell(ws.cell(row=row, column=4), sig.get("type"))
                base_cell(ws.cell(row=row, column=5), sig.get("specifics"))
                base_cell(ws.cell(row=row, column=6), sig.get("date"), center=True)
                write_link(ws.cell(row=row, column=7), sig.get("source_url"), "Source ↗")
                base_cell(ws.cell(row=row, column=8), c.get("role"), center=True)
                nm = c.get("name", "—")
                if c.get("title"):
                    nm = f"{nm} — {c['title']}"
                base_cell(ws.cell(row=row, column=9), nm)
                write_link(ws.cell(row=row, column=10), c.get("linkedin"), "LinkedIn ↗")
                base_cell(ws.cell(row=row, column=11), "; ".join(c.get("person_findings", []) or []))
                base_cell(ws.cell(row=row, column=12), c.get("pitch", ""))
                write_link(ws.cell(row=row, column=13), comp.get("company_linkedin"), "Company ↗")
                write_link(ws.cell(row=row, column=14), comp.get("website"),
                           (comp.get("website") or "").replace("https://", "").rstrip("/") or "Not found")
                base_cell(ws.cell(row=row, column=15), comp.get("hq"))
                base_cell(ws.cell(row=row, column=16), comp.get("public_email"))
                base_cell(ws.cell(row=row, column=17), comp.get("public_phone"))
                base_cell(ws.cell(row=row, column=18), notes, caution=caution)
                row += 1

        if band % 2 == 0:
            for r in rows_here:
                for col in range(1, ncols + 1):
                    ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=stripe)
        for r in rows_here:
            for col in range(1, ncols + 1):
                ws.cell(row=r, column=col).border = BORDER
        band += 1

    ws.freeze_panes = "C4"
    ws.auto_filter.ref = f"A{HDR}:{last}{row - 1}"

    ws2 = wb.create_sheet("Sources & Method")
    ws2.column_dimensions["A"].width = 118
    lines = [
        ("Methodology", True),
        (f"Signal criteria: {cfg.get('signal_criteria', '')}. Compiled {cfg.get('compiled_date', '')}.", False),
        ("Leadership names verified against company-owned pages (leadership / board / IR) and "
         "recent press wherever possible. LinkedIn URLs included only where the profile slug "
         "surfaced verbatim in public sources — none guessed. 'Not found' means exactly that.", False),
        ("No personal emails or mobile numbers included; only company-published channels. "
         "Enrich blanks via your data platform.", False),
        ("", False),
        ("Market-level sources", True),
        *[(u, False) for u in cfg.get("market_sources", [])],
        ("", False),
        ("Per-company sources", True),
        *[(f"{comp.get('company')}: " + " | ".join(comp.get("sources", [])[:4]), False) for comp in data],
    ]
    for i, (txt, bold) in enumerate(lines, 1):
        c = ws2.cell(row=i, column=1, value=txt)
        c.font = Font(name="Arial", size=11 if bold else 10, bold=bold,
                      color=cfg.get("brand_color", "E8611A") if bold else "000000")
        c.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(out_path)
    contacts_n = sum(len(c.get("contacts", [])) for c in data)
    print(json.dumps({"saved": out_path, "companies": len(data), "contacts": contacts_n,
                      "layout": layout}))


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--data", required=True)
    p.add_argument("--config", required=True)
    p.add_argument("--out", required=True)
    a = p.parse_args()
    with open(a.data) as f:
        data = json.load(f)
    with open(a.config) as f:
        cfg = json.load(f)
    build(data, cfg, a.out)
